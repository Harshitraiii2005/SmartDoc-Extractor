from flask import Flask, render_template, request, send_file
import os
import spacy
import pdfplumber
import pytesseract
import pandas as pd
from docx import Document
from PIL import Image
from fpdf import FPDF
import tempfile

# ✅ Tesseract OCR Path
pytesseract.pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"

# ✅ Flask App
app = Flask(__name__, static_folder="static", template_folder="templates")

# ✅ Load spaCy NER Model
nlp = spacy.load("trained_invoice_ner")

# ✅ PDF Class
class PDF(FPDF):
    def __init__(self):
        super().__init__()
        font_dir = os.path.join("fonts")
        self.add_font("DejaVu", "", os.path.join(font_dir, "DejaVuSans.ttf"), uni=True)
        self.add_font("DejaVu", "B", os.path.join(font_dir, "DejaVuSans-Bold.ttf"), uni=True)
        self.set_font("DejaVu", "", 12)

    def header(self):
        self.set_font("DejaVu", "B", 14)
        self.cell(0, 10, "SmartDoc Extracted Report", ln=True, align="C")
        self.ln(5)

# ✅ Routes
@app.route('/')
def home():
    return render_template("index.html")

@app.route('/app')
def extractor():
    return render_template("app.html")

@app.route('/extract', methods=["POST"])
def extract():
    file = request.files['file']
    content = ""
    mime_type = file.content_type

    # ✅ Extract content
    if mime_type == "text/plain":
        content = file.read().decode("utf-8")
    elif mime_type == "application/pdf":
        with pdfplumber.open(file) as pdf:
            content = "\n".join(page.extract_text() or "" for page in pdf.pages)
    elif mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        doc = Document(file)
        content = "\n".join([p.text for p in doc.paragraphs])
    elif mime_type == "text/csv":
        df = pd.read_csv(file)
        content = df.to_string()
    elif "image" in mime_type:
        image = Image.open(file)
        content = pytesseract.image_to_string(image)

    # ✅ NER
    doc_nlp = nlp(content)
    entities = [{"text": ent.text.strip(), "label": ent.label_} for ent in doc_nlp.ents]

    # ✅ Generate PDF
    pdf = PDF()
    pdf.add_page()

    # -- Extracted Text --
    pdf.set_font("DejaVu", "B", 14)
    pdf.cell(0, 10, "Extracted Text", ln=True)
    pdf.set_font("DejaVu", "", 11)
    pdf.multi_cell(0, 8, content)
    pdf.ln(5)

    # -- Named Entities Table --
    pdf.set_font("DejaVu", "B", 14)
    pdf.cell(0, 10, "Named Entities", ln=True)
    pdf.ln(3)

    if entities:
        col1_width = 50
        col2_width = 130
        pdf.set_font("DejaVu", "B", 12)
        pdf.cell(col1_width, 8, "Label", border=1)
        pdf.cell(col2_width, 8, "Value", border=1)
        pdf.ln()

        pdf.set_font("DejaVu", "", 11)
        for ent in entities:
            label = ent['label'].strip()
            text = ent['text'].strip()
            pdf.cell(col1_width, 8, label, border=1)
            pdf.multi_cell(col2_width, 8, text, border=1)
    else:
        pdf.set_font("DejaVu", "", 11)
        pdf.cell(0, 8, "No named entities found.", ln=True)

    # ✅ Save PDF
    temp_pdf_path = os.path.join(tempfile.gettempdir(), "output.pdf")
    pdf.output(temp_pdf_path)

    return render_template("app.html", extracted_text=content, entities=entities, pdf_path="/download_pdf")

@app.route('/download_pdf')
def download_pdf():
    temp_pdf_path = os.path.join(tempfile.gettempdir(), "output.pdf")
    return send_file(temp_pdf_path, as_attachment=True, download_name="SmartDoc_Extracted.pdf")

# ✅ Run App
if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5000, debug=True)
